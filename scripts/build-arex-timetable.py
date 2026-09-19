"""
공항철도(AREX, 서울 ~ 인천공항2터미널 14개 역) 열차시간표 컴파일 스크립트

공항철도 공식 열차시각표(data/공항철도_열차시각표_251229.xlsx)를 파싱하여
초경량 Gzip 인메모리 압축 데이터베이스(arex_subway_timetables.json.gz)로 빌드합니다.

실행: python scripts/build-arex-timetable.py
"""

import os
import sys
import json
import gzip
import datetime
import openpyxl

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if sys.stderr and hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

# 14개 공식 여객 역사 (서울 ~ 인천공항2터미널)
AREX_STATIONS = [
    '서울',
    '공덕',
    '홍대입구',
    '디지털미디어시티',
    '마곡나루',
    '김포공항',
    '계양',
    '검암',
    '청라국제도시',
    '영종',
    '운서',
    '공항화물청사',
    '인천공항1터미널',
    '인천공항2터미널',
]

# 하행: 서울 -> 인천공항2터미널 행 매핑 (도착행, 출발행)
DOWN_STATION_ROWS = [
    ('서울', 13, 14),
    ('공덕', 15, 16),
    ('홍대입구', 17, 18),
    ('디지털미디어시티', 19, 20),
    ('마곡나루', 23, 24),
    ('김포공항', 25, 26),
    ('계양', 27, 28),
    ('검암', 29, 30),
    ('청라국제도시', 31, 32),
    ('영종', 33, 34),
    ('운서', 35, 36),
    ('공항화물청사', 37, 38),
    ('인천공항1터미널', 39, 40),
    ('인천공항2터미널', 41, 42),
]

# 상행: 인천공항2터미널 -> 서울 행 매핑 (도착행, 출발행)
UP_STATION_ROWS = [
    ('인천공항2터미널', 51, 52),
    ('인천공항1터미널', 53, 54),
    ('공항화물청사', 55, 56),
    ('운서', 57, 58),
    ('영종', 59, 60),
    ('청라국제도시', 61, 62),
    ('검암', 63, 64),
    ('계양', 65, 66),
    ('김포공항', 67, 68),
    ('마곡나루', 69, 70),
    ('디지털미디어시티', 73, 74),
    ('홍대입구', 75, 76),
    ('공덕', 77, 78),
    ('서울', 79, 80),
]

def format_cell_time(val) -> str:
    """셀 값을 'HH:mm' 형식의 문자열로 변환합니다."""
    if val is None:
        return ''
    if isinstance(val, (datetime.time, datetime.datetime)):
        return f"{val.hour:02d}:{val.minute:02d}"
    s = str(val).strip()
    if not s or s == '---':
        return ''
    parts = s.split(':')
    if len(parts) >= 2:
        try:
            h = int(parts[0])
            m = int(parts[1])
            return f"{h:02d}:{m:02d}"
        except ValueError:
            return ''
    return ''

def op_key(it):
    """04:00 기준 영업일 순서 키"""
    hh, mm = map(int, it['t'].split(':'))
    if hh < 4:
        hh += 24
    return hh * 60 + mm

def build_arex_timetable():
    excel_path = os.path.join(os.getcwd(), 'data', '공항철도_열차시각표_251229.xlsx')
    if not os.path.exists(excel_path):
        print(f"오류: 시각표 파일 없음: {excel_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    station_db = {stn: {} for stn in AREX_STATIONS}

    sheet_configs = [
        ('평일', 'DAY'),
        ('휴일', 'END'),
    ]

    for sheet_name, tag in sheet_configs:
        if sheet_name not in wb.sheetnames:
            print(f"경고: 시트 {sheet_name} 없음")
            continue

        ws = wb[sheet_name]
        max_col = ws.max_column

        # 1. 하행 (서울 -> 인천공항2터미널) 파싱
        # 열차번호 Row 4, 종착역 Row 7, 비고 Row 43
        down_train_count = 0
        for c in range(2, max_col + 1):
            train_no_cell = ws.cell(4, c).value
            if not train_no_cell:
                continue
            train_no = str(train_no_cell).strip()
            if not train_no:
                continue

            dest_cell = ws.cell(7, c).value or '인천공항2터미널'
            dest = str(dest_cell).strip()
            bigo_cell = ws.cell(43, c).value or ''
            is_express = 1 if (str(bigo_cell).strip() == '직통' or train_no.startswith('A1')) else 0

            for stn_name, arr_r, dep_r in DOWN_STATION_ROWS:
                arr_val = ws.cell(arr_r, c).value
                dep_val = ws.cell(dep_r, c).value

                # 무정차 통과역('---')인 경우 스킵
                if str(arr_val).strip() == '---' or str(dep_val).strip() == '---':
                    continue

                time_str = format_cell_time(dep_val)
                # 출발 시간이 없고 종착역인 경우 도착 시간 사용
                if not time_str and stn_name == dest:
                    time_str = format_cell_time(arr_val)

                if not time_str:
                    continue

                gkey = f"공항철도_{tag}_DOWN"
                if gkey not in station_db[stn_name]:
                    station_db[stn_name][gkey] = []

                item = {
                    'no': train_no,
                    't': time_str,
                    'd': dest,
                    'e': is_express
                }
                station_db[stn_name][gkey].append(item)

            down_train_count += 1

        # 2. 상행 (인천공항2터미널 -> 서울) 파싱
        # 열차번호 Row 46, 종착역 Row 49, 비고 Row 85
        up_train_count = 0
        for c in range(2, max_col + 1):
            train_no_cell = ws.cell(46, c).value
            if not train_no_cell:
                continue
            train_no = str(train_no_cell).strip()
            if not train_no:
                continue

            dest_cell = ws.cell(49, c).value or '서울'
            dest = str(dest_cell).strip()
            bigo_cell = ws.cell(85, c).value or ''
            is_express = 1 if (str(bigo_cell).strip() == '직통' or train_no.startswith('A1')) else 0

            for stn_name, arr_r, dep_r in UP_STATION_ROWS:
                arr_val = ws.cell(arr_r, c).value
                dep_val = ws.cell(dep_r, c).value

                if str(arr_val).strip() == '---' or str(dep_val).strip() == '---':
                    continue

                time_str = format_cell_time(dep_val)
                if not time_str and stn_name == dest:
                    time_str = format_cell_time(arr_val)

                if not time_str:
                    continue

                gkey = f"공항철도_{tag}_UP"
                if gkey not in station_db[stn_name]:
                    station_db[stn_name][gkey] = []

                item = {
                    'no': train_no,
                    't': time_str,
                    'd': dest,
                    'e': is_express
                }
                station_db[stn_name][gkey].append(item)

            up_train_count += 1

        print(f"[{sheet_name}] 하행: {down_train_count}편, 상행: {up_train_count}편 파싱 완료")

    # 3. 토요일(SAT)은 공휴일(END)과 동일 적용
    for stn_name in AREX_STATIONS:
        station_db[stn_name]['공항철도_SAT_UP'] = station_db[stn_name].get('공항철도_END_UP', [])[:]
        station_db[stn_name]['공항철도_SAT_DOWN'] = station_db[stn_name].get('공항철도_END_DOWN', [])[:]

    # 4. 시간순 정렬 및 중복 제거
    total_entries = 0
    for stn_name in AREX_STATIONS:
        for gkey in list(station_db[stn_name].keys()):
            items = station_db[stn_name][gkey]
            # 영업시간 04:00 기준 정렬
            sorted_items = sorted(items, key=op_key)
            # 중복 제거 (no, t, d 동일한 경우)
            unique = []
            seen = set()
            for it in sorted_items:
                k = (it['no'], it['t'], it['d'])
                if k not in seen:
                    seen.add(k)
                    unique.append(it)
            station_db[stn_name][gkey] = unique
            total_entries += len(unique)

    # 메타데이터 구성
    final_payload = {
        'meta': {
            'line': '공항철도',
            'subwayId': '1065',
            'totalStations': len(AREX_STATIONS),
            'totalTimetableEntries': total_entries,
            'stations': AREX_STATIONS,
            'updatedAt': datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'source': '공항철도 공식 열차시각표 (공항철도_열차시각표_251229.xlsx)',
        },
        'data': station_db
    }

    out_dir = os.path.join(os.getcwd(), 'src', 'data', 'subway', 'timetables')
    os.makedirs(out_dir, exist_ok=True)
    out_gz = os.path.join(out_dir, 'arex_subway_timetables.json.gz')

    json_bytes = json.dumps(final_payload, ensure_ascii=False, separators=(',', ':')).encode('utf-8')
    with gzip.open(out_gz, 'wb', compresslevel=9) as f:
        f.write(json_bytes)

    gz_size = os.path.getsize(out_gz)
    print(f"\n빌드 성공!")
    print(f"- 대상 파일: {out_gz}")
    print(f"- 압축 크기: {gz_size / 1024:.1f} KB")
    print(f"- 총 역사 수: {len(AREX_STATIONS)}개")
    print(f"- 총 시각표 레코드 수: {total_entries:,}건")

if __name__ == '__main__':
    build_arex_timetable()
