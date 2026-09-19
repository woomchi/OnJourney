"""
경의중앙선(문산~용산~청량리~용문~지평 53개 역) 열차시각표 컴파일 스크립트

엑셀 시각표(data/경의중앙선_열차시간표.xlsx)를 파싱하여
고효율 Gzip 인메모리 압축 데이터베이스(gyeongui_subway_timetables.json.gz)로 빌드합니다.

실행: python scripts/build-gyeongui-timetable.py
"""

import os
import sys
import json
import gzip
import datetime
import openpyxl

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if sys.stderr and hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

# 표준 역명 매핑 (엑셀 축약어 -> 표준 역명)
STATION_NAME_MAP = {
    '1양정': '양정',
    '양정': '양정',
    '1양원': '양원',
    '양원': '양원',
    '효창공': '효창공원앞',
    '효창공원앞': '효창공원앞',
    '홍대입': '홍대입구',
    '홍대입구': '홍대입구',
    '디엠시': '디지털미디어시티',
    '디지털미디어시티': '디지털미디어시티',
    '항공대': '한국항공대',
    '한국항공대': '한국항공대',
    '지평': '지평',
    '용문': '용문',
    '원덕': '원덕',
    '양평': '양평',
    '오빈': '오빈',
    '아신': '아신',
    '국수': '국수',
    '신원': '신원',
    '양수': '양수',
    '운길산': '운길산',
    '팔당': '팔당',
    '도심': '도심',
    '덕소': '덕소',
    '도농': '도농',
    '구리': '구리',
    '망우': '망우',
    '상봉': '상봉',
    '중랑': '중랑',
    '회기': '회기',
    '청량리': '청량리',
    '왕십리': '왕십리',
    '응봉': '응봉',
    '옥수': '옥수',
    '한남': '한남',
    '서빙고': '서빙고',
    '이촌': '이촌',
    '용산': '용산',
    '공덕': '공덕',
    '서강대': '서강대',
    '가좌': '가좌',
    '수색': '수색',
    '강매': '강매',
    '행신': '행신',
    '능곡': '능곡',
    '대곡': '대곡',
    '곡산': '곡산',
    '백마': '백마',
    '풍산': '풍산',
    '일산': '일산',
    '탄현': '탄현',
    '야당': '야당',
    '운정': '운정',
    '금릉': '금릉',
    '금촌': '금촌',
    '월롱': '월롱',
    '파주': '파주',
    '문산': '문산',
}

def normalize_station_name(raw: str) -> str:
    if not raw:
        return ''
    clean = str(raw).strip()
    return STATION_NAME_MAP.get(clean, clean)

def time_to_hhmm(val) -> str:
    if val is None:
        return ''
    if isinstance(val, datetime.time):
        return f"{val.hour:02d}:{val.minute:02d}"
    if isinstance(val, datetime.datetime):
        return f"{val.hour:02d}:{val.minute:02d}"
    s = str(val).strip()
    if not s:
        return ''
    parts = s.split(':')
    if len(parts) >= 2:
        try:
            h = int(parts[0])
            m = int(parts[1])
            return f"{h:02d}:{m:02d}"
        except ValueError:
            return ''
    return ''

def is_express_train(train_no: str) -> int:
    """K57xx, K58xx 급행 열차 식별"""
    t = str(train_no).strip().upper()
    if t.startswith('K57') or t.startswith('K58'):
        return 1
    return 0

def to_operational_minutes(time_str: str) -> int:
    if not time_str:
        return 0
    parts = time_str.split(':')
    h = int(parts[0])
    m = int(parts[1]) if len(parts) > 1 else 0
    adjusted_h = h + 24 if h < 5 else h
    return adjusted_h * 60 + m

def process_sheet(ws, day_type: str, direction: str, database: dict, all_stations_set: set):
    """
    각 시트를 읽어 database[stn][key] 리스트에 추가합니다.
    day_type: 'DAY' (평일) 또는 'END' (휴일/토/일)
    direction: 'UP' (상행/문산방면) 또는 'DOWN' (하행/용문·지평방면)
    """
    group_key = f"경의중앙선_{day_type}_{direction}"

    # 1. Col 1에서 역명 행과 출발시각 행 매핑 탐색 (Row 4부터)
    station_rows = []
    for r in range(4, ws.max_row + 1):
        val = ws.cell(r, 1).value
        if val and str(val).strip():
            clean_str = str(val).strip()
            if clean_str in ('연계열번', '비고', '열차번호', '시발역', '종착역'):
                continue
            stn = normalize_station_name(clean_str)
            station_rows.append((r, stn))
            all_stations_set.add(stn)
            if stn not in database:
                database[stn] = {}
            if group_key not in database[stn]:
                database[stn][group_key] = []

    # 2. 열차별(Col 2부터) 순회
    train_count = 0
    for c in range(2, ws.max_column + 1):
        dest_stn_raw = ws.cell(2, c).value
        train_no_raw = ws.cell(3, c).value

        if not train_no_raw:
            continue

        train_no = str(train_no_raw).strip()
        dest_stn = normalize_station_name(str(dest_stn_raw)) if dest_stn_raw else ''
        is_express = is_express_train(train_no)
        train_count += 1

        for idx, (r, stn) in enumerate(station_rows):
            # 다음 역 행 번호
            next_r = station_rows[idx + 1][0] if idx + 1 < len(station_rows) else ws.max_row + 1

            times = []
            for sub_r in range(r, next_r):
                cell_val = ws.cell(sub_r, c).value
                t_str = time_to_hhmm(cell_val)
                if t_str:
                    times.append(t_str)

            if not times:
                continue

            # 승차 기준은 마지막 시각(출발시각), 단 종착역이면 도착시각
            effective_time = times[-1]

            database[stn][group_key].append({
                'no': train_no,
                't': effective_time,
                'd': dest_stn,
                'e': is_express,
            })

    print(f"[{ws.title}] {train_count} trains processed for key: {group_key}")

def main():
    root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    excel_path = os.path.join(root_dir, 'data', '경의중앙선_열차시간표.xlsx')
    output_dir = os.path.join(root_dir, 'src', 'data', 'subway', 'timetables')
    gz_output_path = os.path.join(output_dir, 'gyeongui_subway_timetables.json.gz')
    json_output_path = os.path.join(output_dir, 'gyeongui_subway_timetables.json')

    print(f"Reading Excel: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    database = {}
    all_stations_set = set()

    # 시트 매핑 (공백 및 괄호 대응)
    sheet_configs = [
        (['경의중앙선 상행(평일)', '경의중앙선상행(평일)'], 'DAY', 'UP'),
        (['경의중앙선 하행(평일)', '경의중앙선하행(평일)'], 'DAY', 'DOWN'),
        (['경의중앙선 상행(휴일)', '경의중앙선상행(휴일)'], 'END', 'UP'),
        (['경의중앙선 하행(휴일)', '경의중앙선하행(휴일)'], 'END', 'DOWN'),
    ]

    for sheet_names, day_type, direction in sheet_configs:
        found_ws = None
        for sname in sheet_names:
            if sname in wb.sheetnames:
                found_ws = wb[sname]
                break
        if found_ws:
            process_sheet(found_ws, day_type, direction, database, all_stations_set)
        else:
            print(f"WARNING: Sheet not found for: {sheet_names}")

    # 토요일(SAT) 키를 휴일(END)과 동일하게 복제
    for stn, group in database.items():
        if '경의중앙선_END_UP' in group and '경의중앙선_SAT_UP' not in group:
            group['경의중앙선_SAT_UP'] = list(group['경의중앙선_END_UP'])
        if '경의중앙선_END_DOWN' in group and '경의중앙선_SAT_DOWN' not in group:
            group['경의중앙선_SAT_DOWN'] = list(group['경의중앙선_END_DOWN'])

    # 운행 시각 기준 정렬 및 중복 제거
    total_entries = 0
    for stn, group in database.items():
        for gkey, items in group.items():
            unique_dict = {}
            for item in items:
                k = f"{item['no']}_{item['t']}"
                if k not in unique_dict:
                    unique_dict[k] = item
            sorted_items = sorted(
                unique_dict.values(),
                key=lambda x: to_operational_minutes(x['t'])
            )
            group[gkey] = sorted_items
            total_entries += len(sorted_items)

    final_db = {
        'meta': {
            'updatedAt': datetime.datetime.now().isoformat(),
            'source': '한국철도공사 경의중앙선 열차시간표',
            'line': '경의중앙선',
            'subwayId': '1063',
            'totalStations': len(all_stations_set),
            'totalTimetableEntries': total_entries,
            'stations': sorted(list(all_stations_set)),
        },
        'data': database,
    }

    print(f"\nTotal Stations: {len(all_stations_set)}")
    print(f"Total Timetable Entries: {total_entries}")

    # JSON 저장
    os.makedirs(output_dir, exist_ok=True)
    json_bytes = json.dumps(final_db, ensure_ascii=False, indent=2).encode('utf-8')
    with open(json_output_path, 'wb') as f:
        f.write(json_bytes)
    print(f"JSON saved: {json_output_path} ({len(json_bytes) / 1024:.1f} KB)")

    # GZIP 압축 저장
    gz_bytes = gzip.compress(json_bytes, compresslevel=9)
    with open(gz_output_path, 'wb') as f:
        f.write(gz_bytes)
    print(f"GZ saved: {gz_output_path} ({len(gz_bytes) / 1024:.1f} KB)")

if __name__ == '__main__':
    main()
