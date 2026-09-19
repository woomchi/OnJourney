"""
수인분당선(청량리/왕십리~수원~인천 63개 역) 열차시각표 컴파일 스크립트

엑셀 시각표(data/수인분당선_열차시간표.xlsx)를 파싱하여
고효율 Gzip 인메모리 압축 데이터베이스(suinbundang_subway_timetables.json.gz)로 빌드합니다.

실행: python scripts/build-suinbundang-timetable.py
"""

import os
import sys
import json
import gzip
import datetime
import openpyxl

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if sys.stderr and hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

# 표준 역명 매핑 (엑셀 축약어 -> 표준 역명)
STATION_NAME_MAP = {
    '신인천': '인천',
    '인천': '인천',
    '신포': '신포',
    '숭의': '숭의',
    '인하대': '인하대',
    '송도': '송도',
    '연수': '연수',
    '원인재': '원인재',
    '남동인': '남동인더스파크',
    '남동인더스파크': '남동인더스파크',
    '호구포': '호구포',
    '인천논': '인천논현',
    '인천논현': '인천논현',
    '소래포': '소래포구',
    '소래포구': '소래포구',
    '월곶': '월곶',
    '달월': '달월',
    '오이도': '오이도',
    '정왕': '정왕',
    '신길온': '신길온천',
    '신길온천': '신길온천',
    '안산': '안산',
    '초지': '초지',
    '고잔': '고잔',
    '중앙': '중앙',
    '한대앞': '한대앞',
    '사리': '사리',
    '야목': '야목',
    '어천': '어천',
    '오목천': '오목천',
    '고색': '고색',
    '신수원': '수원',
    '수원': '수원',
    '매교': '매교',
    '수원시': '수원시청',
    '수원시청': '수원시청',
    '매탄권': '매탄권선',
    '매탄권선': '매탄권선',
    '망포': '망포',
    '영통': '영통',
    '청명': '청명',
    '상갈': '상갈',
    '기흥': '기흥',
    '신갈': '신갈',
    '구성': '구성',
    '보정': '보정',
    '죽전': '죽전',
    '오리': '오리',
    '미금': '미금',
    '정자': '정자',
    '수내': '수내',
    '서현': '서현',
    '이매': '이매',
    '야탑': '야탑',
    '모란': '모란',
    '태평': '태평',
    '가천대': '가천대',
    '복정': '복정',
    '수서': '수서',
    '대모산': '대모산입구',
    '대모산입구': '대모산입구',
    '개포동': '개포동',
    '구룡역': '구룡',
    '구룡': '구룡',
    '도곡': '도곡',
    '한티': '한티',
    '선릉': '선릉',
    '선정릉': '선정릉',
    '강남구': '강남구청',
    '강남구청': '강남구청',
    '로데오': '압구정로데오',
    '압구정로데오': '압구정로데오',
    '서울숲': '서울숲',
    '왕십리': '왕십리',
    '청량리': '청량리',
}

def normalize_station_name(raw: str) -> str:
    if not raw:
        return ''
    clean = str(raw).strip()
    return STATION_NAME_MAP.get(clean, clean)

def time_to_hhmm(val) -> str:
    if val is None:
        return ''
    if isinstance(val, datetime.time):
        return f"{val.hour:02d}:{val.minute:02d}"
    if isinstance(val, datetime.datetime):
        return f"{val.hour:02d}:{val.minute:02d}"
    s = str(val).strip()
    if not s:
        return ''
    parts = s.split(':')
    if len(parts) >= 2:
        try:
            h = int(parts[0])
            m = int(parts[1])
            return f"{h:02d}:{m:02d}"
        except ValueError:
            return ''
    return ''

def is_express_train(train_no: str) -> int:
    """K63xx(분당선 급행), K66xx(수인선 급행) 등 판별"""
    t = str(train_no).strip().upper()
    if t.startswith('K63') or t.startswith('K66'):
        return 1
    return 0

def to_operational_minutes(time_str: str) -> int:
    if not time_str:
        return 0
    parts = time_str.split(':')
    h = int(parts[0])
    m = int(parts[1]) if len(parts) > 1 else 0
    adjusted_h = h + 24 if h < 5 else h
    return adjusted_h * 60 + m

def process_sheet(ws, day_type: str, direction: str, database: dict, all_stations_set: set):
    """
    각 시트를 읽어 database[stn][key] 리스트에 추가합니다.
    day_type: 'DAY' (평일) 또는 'END' (휴일/토/일)
    direction: 'UP' (상행/청량리·왕십리방면) 또는 'DOWN' (하행/수원·인천방면)
    """
    group_key = f"수인분당선_{day_type}_{direction}"
    
    # 1. Col 1에서 역 목록 추출 (Row 5부터 2행씩)
    # 홀수 행: 역명 + 도착시각, 짝수 행: None + 출발시각
    row_station_map = {}
    for r in range(5, ws.max_row + 1, 2):
        cell_val = ws.cell(r, 1).value
        if cell_val:
            stn = normalize_station_name(str(cell_val))
            row_station_map[r] = stn
            all_stations_set.add(stn)
            if stn not in database:
                database[stn] = {}
            if group_key not in database[stn]:
                database[stn][group_key] = []

    # 2. 각 열차(열 Col 2부터) 순회
    train_count = 0
    for c in range(2, ws.max_column + 1):
        start_stn_raw = ws.cell(2, c).value
        dest_stn_raw = ws.cell(3, c).value
        train_no_raw = ws.cell(4, c).value

        if not train_no_raw:
            continue

        train_no = str(train_no_raw).strip()
        dest_stn = normalize_station_name(str(dest_stn_raw)) if dest_stn_raw else ''
        is_express = is_express_train(train_no)
        train_count += 1

        for r, stn in row_station_map.items():
            arr_val = ws.cell(r, c).value      # 홀수 행: 도착시각
            dep_val = ws.cell(r + 1, c).value  # 짝수 행: 출발시각

            dep_time = time_to_hhmm(dep_val)
            arr_time = time_to_hhmm(arr_val)

            # 승차 기준은 출발시각 우선 채택. 단, 종착역이거나 출발시각이 없으면 도착시각 사용
            effective_time = dep_time if dep_time else arr_time
            if not effective_time:
                continue

            database[stn][group_key].append({
                'no': train_no,
                't': effective_time,
                'd': dest_stn,
                'e': is_express,
            })

    print(f"[{ws.title}] {train_count} trains processed for key: {group_key}")

def main():
    root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    excel_path = os.path.join(root_dir, 'data', '수인분당선_열차시간표.xlsx')
    output_dir = os.path.join(root_dir, 'src', 'data', 'subway', 'timetables')
    gz_output_path = os.path.join(output_dir, 'suinbundang_subway_timetables.json.gz')
    json_output_path = os.path.join(output_dir, 'suinbundang_subway_timetables.json')

    print(f"Reading Excel: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    database = {}
    all_stations_set = set()

    # 시트 매핑
    sheet_configs = [
        ('평일상행', 'DAY', 'UP'),
        ('평일하행', 'DAY', 'DOWN'),
        ('휴일상행', 'END', 'UP'),
        ('휴일하행', 'END', 'DOWN'),
    ]

    for sheet_name, day_type, direction in sheet_configs:
        if sheet_name in wb.sheetnames:
            process_sheet(wb[sheet_name], day_type, direction, database, all_stations_set)
        else:
            print(f"WARNING: Sheet '{sheet_name}' not found in workbook!")

    # 토요일(SAT) 키를 휴일(END)과 동일하게 복제하여 일관된 룩업 보장
    for stn, group in database.items():
        if '수인분당선_END_UP' in group and '수인분당선_SAT_UP' not in group:
            group['수인분당선_SAT_UP'] = list(group['수인분당선_END_UP'])
        if '수인분당선_END_DOWN' in group and '수인분당선_SAT_DOWN' not in group:
            group['수인분당선_SAT_DOWN'] = list(group['수인분당선_END_DOWN'])

    # 운행 시각 기준 정렬 및 중복 제거
    total_entries = 0
    for stn, group in database.items():
        for gkey, items in group.items():
            # 중복 제거 (no + t 기준)
            unique_dict = {}
            for item in items:
                k = f"{item['no']}_{item['t']}"
                if k not in unique_dict:
                    unique_dict[k] = item
            sorted_items = sorted(
                unique_dict.values(),
                key=lambda x: to_operational_minutes(x['t'])
            )
            group[gkey] = sorted_items
            total_entries += len(sorted_items)

    final_db = {
        'meta': {
            'updatedAt': datetime.datetime.now().isoformat(),
            'source': '한국철도공사 수인분당선 열차시간표',
            'line': '수인분당선',
            'subwayId': '1075',
            'totalStations': len(all_stations_set),
            'totalTimetableEntries': total_entries,
            'stations': sorted(list(all_stations_set)),
        },
        'data': database,
    }

    print(f"\nTotal Stations: {len(all_stations_set)}")
    print(f"Total Timetable Entries: {total_entries}")

    # JSON 저장
    os.makedirs(output_dir, exist_ok=True)
    json_bytes = json.dumps(final_db, ensure_ascii=False, indent=2).encode('utf-8')
    with open(json_output_path, 'wb') as f:
        f.write(json_bytes)
    print(f"JSON saved: {json_output_path} ({len(json_bytes) / 1024:.1f} KB)")

    # GZIP 압축 저장
    gz_bytes = gzip.compress(json_bytes, compresslevel=9)
    with open(gz_output_path, 'wb') as f:
        f.write(gz_bytes)
    print(f"GZ saved: {gz_output_path} ({len(gz_bytes) / 1024:.1f} KB)")

if __name__ == '__main__':
    main()
