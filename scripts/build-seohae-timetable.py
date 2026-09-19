"""
서해선(원시~대곡~일산 21개 역) 열차시간표 컴파일 스크립트

엑셀 시각표(data/서해선_열차시간표.xlsx)를 파싱하여
초경량 Gzip 인메모리 압축 데이터베이스(seohae_subway_timetables.json.gz)로 빌드합니다.

실행: python scripts/build-seohae-timetable.py
"""

import os
import sys
import json
import gzip
import datetime
import openpyxl

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if sys.stderr and hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

# 표준 역명 매핑
STATION_NAME_MAP = {
    '원시': '원시',
    '시우': '시우',
    '신초지': '초지',
    '초지': '초지',
    '선부': '선부',
    '달미': '달미',
    '시흥능': '시흥능곡',
    '시흥능곡': '시흥능곡',
    '시흥청': '시흥시청',
    '시흥시청': '시흥시청',
    '신신현': '신현',
    '신현': '신현',
    '신신천': '신천',
    '신천': '신천',
    '시흥대': '시흥대야',
    '시흥대야': '시흥대야',
    '소새울': '소새울',
    '신소사': '소사',
    '소사': '소사',
    '부천종': '부천종합운동장',
    '부천종합운동장': '부천종합운동장',
    '원종': '원종',
    '신김포': '김포공항',
    '김포공항': '김포공항',
    '능곡': '능곡',
    '대곡': '대곡',
    '곡산': '곡산',
    '백마': '백마',
    '풍산': '풍산',
    '일산': '일산',
}

# 21개 공식 역사 순서 (원시 ~ 일산)
SEOHAE_STATIONS = [
    '원시',
    '시우',
    '초지',
    '선부',
    '달미',
    '시흥능곡',
    '시흥시청',
    '신현',
    '신천',
    '시흥대야',
    '소새울',
    '소사',
    '부천종합운동장',
    '원종',
    '김포공항',
    '능곡',
    '대곡',
    '곡산',
    '백마',
    '풍산',
    '일산',
]

def format_cell_time(val) -> str:
    """셀 값을 'HH:mm' 형식의 문자열로 변환합니다."""
    if val is None:
        return ''
    if isinstance(val, datetime.time):
        return f"{val.hour:02d}:{val.minute:02d}"
    if isinstance(val, datetime.datetime):
        return f"{val.hour:02d}:{val.minute:02d}"
    s = str(val).strip()
    if not s:
        return ''
    parts = s.split(':')
    if len(parts) >= 2:
        try:
            h = int(parts[0])
            m = int(parts[1])
            return f"{h:02d}:{m:02d}"
        except ValueError:
            return ''
    return ''

def parse_seohae_sheet(ws, direction: str, week_tag: str, station_db: dict):
    """
    서해선 단일 시트 파싱
    - direction: 'UP' 또는 'DOWN'
    - week_tag: 'DAY', 'SAT', 'END'
    """
    max_col = ws.max_column
    max_row = ws.max_row

    # 1. 역사 행 위치 식별
    # 서해선 엑셀은 역별로 짝수 행에 도착/역명, 홀수 행에 출발이 위치함
    station_rows = [] # list of (station_name, arr_row, dep_row)

    if direction == 'UP':
        # 상행: 원시(Row 4) ~ 일산(Row 44)
        for r in range(4, min(46, max_row + 1), 2):
            raw_stn = ws.cell(r, 1).value
            if raw_stn and str(raw_stn).strip():
                clean_name = str(raw_stn).strip()
                official = STATION_NAME_MAP.get(clean_name)
                if official and official in SEOHAE_STATIONS:
                    station_rows.append((official, r, r + 1))
    else:
        # 하행: 일산(Row 28) ~ 원시(Row 68)
        for r in range(28, min(70, max_row + 1), 2):
            raw_stn = ws.cell(r, 1).value
            if raw_stn and str(raw_stn).strip():
                clean_name = str(raw_stn).strip()
                official = STATION_NAME_MAP.get(clean_name)
                if official and official in SEOHAE_STATIONS:
                    station_rows.append((official, r, r + 1))

    print(f"[{ws.title}] 인식된 역사 수: {len(station_rows)}개")

    # 2. 열(Col 2 ~ max_col) 순회하며 열차 파싱
    train_count = 0
    for c in range(2, max_col + 1):
        train_no_cell = ws.cell(3, c).value
        if not train_no_cell:
            continue
        train_no = str(train_no_cell).strip()
        if not train_no:
            continue

        raw_dest = ws.cell(2, c).value
        dest = STATION_NAME_MAP.get(str(raw_dest).strip(), str(raw_dest).strip()) if raw_dest else (
            '일산' if direction == 'UP' else '원시'
        )

        for stn_name, arr_r, dep_r in station_rows:
            # 출발 시각 우선, 없으면 도착 시각
            dep_val = ws.cell(dep_r, c).value if dep_r <= max_row else None
            arr_val = ws.cell(arr_r, c).value

            time_str = format_cell_time(dep_val)
            if not time_str:
                time_str = format_cell_time(arr_val)

            if not time_str:
                continue

            gkey = f"서해선_{week_tag}_{direction}"
            if gkey not in station_db[stn_name]:
                station_db[stn_name][gkey] = []

            item = {
                'no': train_no,
                't': time_str,
                'd': dest
            }
            station_db[stn_name][gkey].append(item)

        train_count += 1

    print(f"[{ws.title}] 총 파싱 열차 수: {train_count}편")

def build_seohae_timetable():
    excel_path = os.path.join(os.getcwd(), 'data', '서해선_열차시간표.xlsx')
    if not os.path.exists(excel_path):
        print(f"오류: 엑셀 파일이 없습니다: {excel_path}")
        sys.exit(1)

    print(f"=== 서해선(21개 역) 열차시간표 빌드 시작: {excel_path} ===")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    station_db = {stn: {} for stn in SEOHAE_STATIONS}

    # 1. 서해선 상행(평일)
    if '서해선 상행(평일)' in wb.sheetnames:
        parse_seohae_sheet(wb['서해선 상행(평일)'], 'UP', 'DAY', station_db)

    # 2. 서해선 하행(평일)
    if '서해선 하행(평일)' in wb.sheetnames:
        parse_seohae_sheet(wb['서해선 하행(평일)'], 'DOWN', 'DAY', station_db)

    # 3. 서해선 상행(휴일) -> END 및 SAT 공통 적용
    if '서해선 상행(휴일)' in wb.sheetnames:
        parse_seohae_sheet(wb['서해선 상행(휴일)'], 'UP', 'END', station_db)
        parse_seohae_sheet(wb['서해선 상행(휴일)'], 'UP', 'SAT', station_db)

    # 4. 서해선 하행(휴일) -> END 및 SAT 공통 적용
    if '서해선 하행(휴일)' in wb.sheetnames:
        parse_seohae_sheet(wb['서해선 하행(휴일)'], 'DOWN', 'END', station_db)
        parse_seohae_sheet(wb['서해선 하행(휴일)'], 'DOWN', 'SAT', station_db)

    # 3. 정렬 및 중복 제거
    total_entries = 0
    for stn in SEOHAE_STATIONS:
        for gkey in station_db[stn]:
            # 시간순 정렬 (새벽 04:00 기준)
            def op_time(item):
                h, m = map(int, item['t'].split(':'))
                if h < 4:
                    h += 24
                return h * 60 + m

            unique_items = []
            seen = set()
            for it in sorted(station_db[stn][gkey], key=op_time):
                k = (it['no'], it['t'], it['d'])
                if k not in seen:
                    seen.add(k)
                    unique_items.append(it)

            station_db[stn][gkey] = unique_items
            total_entries += len(unique_items)

    final_payload = {
        'meta': {
            'line': '서해선',
            'subwayId': '1093',
            'totalStations': len(SEOHAE_STATIONS),
            'totalTimetableEntries': total_entries,
            'stations': SEOHAE_STATIONS,
            'updatedAt': datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'source': '코레일 서해선 열차시간표 (data/서해선_열차시간표.xlsx)',
        },
        'data': station_db
    }

    out_dir = os.path.join(os.getcwd(), 'src', 'data', 'subway', 'timetables')
    os.makedirs(out_dir, exist_ok=True)
    out_gz = os.path.join(out_dir, 'seohae_subway_timetables.json.gz')

    json_bytes = json.dumps(final_payload, ensure_ascii=False, separators=(',', ':')).encode('utf-8')
    with gzip.open(out_gz, 'wb', compresslevel=9) as f:
        f.write(json_bytes)

    gz_size = os.path.getsize(out_gz)
    print(f"빌드 성공!")
    print(f"- 대상 파일: {out_gz}")
    print(f"- 압축 크기: {gz_size / 1024:.1f} KB")
    print(f"- 총 역사 수: {len(SEOHAE_STATIONS)}개")
    print(f"- 총 시각표 레코드 수: {total_entries:,}건")

if __name__ == '__main__':
    build_seohae_timetable()
