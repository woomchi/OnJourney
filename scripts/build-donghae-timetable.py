"""
동해선(광역전철 부전~태화강) 열차시각표 컴파일 스크립트

엑셀 시각표(data/부산광역시_동해선_열차시각표_20260901.xlsx)를 파싱하여
고효율 Gzip 인메모리 압축 데이터베이스(donghae_subway_timetables.json.gz)로 빌드합니다.

실행: python scripts/build-donghae-timetable.py
"""

import os
import sys
import json
import gzip
import datetime
import openpyxl

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if sys.stderr and hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

# 표준 역명 순서 (부전 -> 태화강 순, 하행 기준)
STATION_ORDER = [
    '부전',
    '거제해맞이',
    '거제',
    '교대',
    '동래',
    '안락',
    '부산원동',
    '재송',
    '센텀',
    '벡스코',
    '신해운대',
    '송정',
    '오시리아',
    '기장',
    '일광',
    '좌천',
    '월내',
    '서생',
    '남창',
    '망양',
    '덕하',
    '개운포',
    '태화강',
]

# 축약명 -> 표준 역명 매핑 사전
STATION_NAME_MAP = {
    '부전': '부전',
    '거제해': '거제해맞이',
    '거제해맞이': '거제해맞이',
    '거제': '거제',
    '부교대': '교대',
    '교대': '교대',
    '동래': '동래',
    '안락': '안락',
    '부산원': '부산원동',
    '부산원동': '부산원동',
    '재송': '재송',
    '센텀': '센텀',
    '벡스코': '벡스코',
    '신해운': '신해운대',
    '신해운대': '신해운대',
    '송정': '송정',
    '오시리': '오시리아',
    '오시리아': '오시리아',
    '기장': '기장',
    '일광': '일광',
    '좌천': '좌천',
    '월내': '월내',
    '서생': '서생',
    '남창': '남창',
    '망양': '망양',
    '덕하': '덕하',
    '개운포': '개운포',
    '태화강': '태화강',
}

def normalize_station_name(raw: str) -> str:
    if not raw:
        return ''
    clean = str(raw).strip()
    return STATION_NAME_MAP.get(clean, clean)

def time_to_hhmm(t: datetime.time) -> str:
    if not t:
        return ''
    return f"{t.hour:02d}:{t.minute:02d}"

def to_operational_minutes(time_str: str) -> int:
    if not time_str:
        return 0
    parts = time_str.split(':')
    h = int(parts[0])
    m = int(parts[1]) if len(parts) > 1 else 0
    adjusted_h = h + 24 if h < 5 else h
    return adjusted_h * 60 + m

def build_donghae_timetables():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    excel_path = os.path.join(base_dir, 'data', '부산광역시_동해선_열차시각표_20260901.xlsx')

    if not os.path.exists(excel_path):
        print(f"Error: 엑셀 파일을 찾을 수 없습니다: {excel_path}", file=sys.stderr)
        sys.exit(1)

    print(f"🚀 동해선 열차시각표 컴파일 시작: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # sheet config: (sheet_name, direction, day_type)
    sheet_configs = [
        ('평일 상행', 'UP', 'WEEKDAY'),
        ('평일 하행', 'DOWN', 'WEEKDAY'),
        ('휴일 상행', 'UP', 'HOLIDAY'),
        ('휴일 하행', 'DOWN', 'HOLIDAY'),
    ]

    # station_timetables[station][key] = [ {no, t, d, e} ]
    station_timetables = {}
    for stn in STATION_ORDER:
        station_timetables[stn] = {}

    total_train_runs = 0

    for sheet_name, direction, day_type in sheet_configs:
        if sheet_name not in wb.sheetnames:
            print(f"Warning: 시트 '{sheet_name}'이 존재하지 않습니다.", file=sys.stderr)
            continue

        sheet = wb[sheet_name]
        key = f"동해선_{direction}_{day_type}"
        print(f"Processing sheet: {sheet_name} -> Key: {key}")

        # 열차 메타 정보 수집 (Row 2: 시발역, Row 3: 종착역, Row 4: 열차번호)
        col_train_meta = {}
        for col in range(2, sheet.max_column + 1):
            train_no = sheet.cell(4, col).value
            if not train_no:
                continue
            start_stn = normalize_station_name(sheet.cell(2, col).value)
            end_stn = normalize_station_name(sheet.cell(3, col).value)
            col_train_meta[col] = {
                'train_no': str(train_no).strip(),
                'start_stn': start_stn,
                'end_stn': end_stn,
            }
            total_train_runs += 1

        # 역별 도착/출발 시간 수집 (Row 5부터 2행씩: r=도착, r+1=출발)
        for r in range(5, sheet.max_row + 1, 2):
            raw_stn = sheet.cell(r, 1).value
            if not raw_stn:
                continue
            stn_name = normalize_station_name(raw_stn)
            if stn_name not in station_timetables:
                station_timetables[stn_name] = {}
            if key not in station_timetables[stn_name]:
                station_timetables[stn_name][key] = []

            for col, meta in col_train_meta.items():
                arr_val = sheet.cell(r, col).value
                dep_val = sheet.cell(r + 1, col).value

                # 탑승객 기준 시간: 출발 시각 우선, 종착역 등 출발 시각이 없으면 도착 시각
                chosen_time = dep_val if isinstance(dep_val, datetime.time) else (
                    arr_val if isinstance(arr_val, datetime.time) else None
                )

                if chosen_time is None:
                    continue  # 해당 역 미정차 열차

                hhmm = time_to_hhmm(chosen_time)
                if not hhmm:
                    continue

                item = {
                    'no': meta['train_no'],
                    't': hhmm,
                    'd': meta['end_stn'],
                    'e': 0,  # 전동열차 일반
                }
                station_timetables[stn_name][key].append(item)

    # 시간순 정렬
    for stn_name, key_dict in station_timetables.items():
        for group_key, items in key_dict.items():
            items.sort(key=lambda x: to_operational_minutes(x['t']))

    out_dir = os.path.join(base_dir, 'src', 'data', 'subway', 'timetables')
    os.makedirs(out_dir, exist_ok=True)
    gz_path = os.path.join(out_dir, 'donghae_subway_timetables.json.gz')
    json_path = os.path.join(out_dir, 'donghae_subway_timetables.json')

    now_str = datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=9))).isoformat()
    unified_db = {
        'meta': {
            'updatedAt': now_str,
            'source': '한국철도공사 동해선 광역전철 열차시각표 (20260901)',
            'totalStations': len(STATION_ORDER),
            'line': '동해선',
            'stations': STATION_ORDER,
        },
        'data': station_timetables,
    }

    # 1. 비압축 JSON 저장
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(unified_db, f, ensure_ascii=False, indent=2)
    json_size_kb = os.path.getsize(json_path) / 1024

    # 2. Gzip 압축 JSON 저장
    with gzip.open(gz_path, 'wt', encoding='utf-8') as f:
        json.dump(unified_db, f, ensure_ascii=False)
    gz_size_kb = os.path.getsize(gz_path) / 1024

    print(f"✅ 동해선 열차시각표 컴파일 완료!")
    print(f"   - 총 정차역: {len(STATION_ORDER)}개 역")
    print(f"   - 총 운행 편수(상/하행, 평/휴일 합계): {total_train_runs}회")
    print(f"   - 비압축 JSON: {json_path} ({json_size_kb:.1f} KB)")
    print(f"   - Gzip 압축 DB: {gz_path} ({gz_size_kb:.1f} KB)")

if __name__ == '__main__':
    build_donghae_timetables()
