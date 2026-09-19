"""
경춘선(청량리/광운대~상봉~춘천 25개 역) 열차시간표 컴파일 스크립트

엑셀 시각표(data/경춘선_열차시간표.xlsx)를 파싱하여
고효율 Gzip 인메모리 압축 데이터베이스(gyeongchun_subway_timetables.json.gz)로 빌드합니다.

실행: python scripts/build-gyeongchun-timetable.py
"""

import os
import sys
import json
import gzip
import datetime
import openpyxl

if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if sys.stderr and hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

# 표준 역명 매핑
STATION_NAME_MAP = {
    '평내호': '평내호평',
    '평내호평': '평내호평',
    '청량리': '청량리',
    '회기': '회기',
    '중랑': '중랑',
    '상봉': '상봉',
    '망우': '망우',
    '신내': '신내',
    '갈매': '갈매',
    '별내': '별내',
    '퇴계원': '퇴계원',
    '사릉': '사릉',
    '금곡': '금곡',
    '천마산': '천마산',
    '마석': '마석',
    '대성리': '대성리',
    '청평': '청평',
    '상천': '상천',
    '가평': '가평',
    '굴봉산': '굴봉산',
    '백양리': '백양리',
    '강촌': '강촌',
    '김유정': '김유정',
    '남춘천': '남춘천',
    '춘천': '춘천',
    '광운대': '광운대',
}

def normalize_station_name(raw: str) -> str:
    if not raw:
        return ''
    clean = str(raw).strip()
    return STATION_NAME_MAP.get(clean, clean)

def time_to_hhmm(val) -> str:
    if val is None:
        return ''
    if isinstance(val, datetime.time):
        return f"{val.hour:02d}:{val.minute:02d}"
    if isinstance(val, datetime.datetime):
        return f"{val.hour:02d}:{val.minute:02d}"
    s = str(val).strip()
    if not s:
        return ''
    parts = s.split(':')
    if len(parts) >= 2:
        try:
            h = int(parts[0])
            m = int(parts[1])
            return f"{h:02d}:{m:02d}"
        except ValueError:
            return ''
    return ''

def is_express_train(train_no: str) -> int:
    """K81xx, K84xx 급행 열차 식별"""
    t = str(train_no).strip().upper()
    if t.startswith('K81') or t.startswith('K84'):
        return 1
    return 0

def to_operational_minutes(time_str: str) -> int:
    if not time_str:
        return 0
    parts = time_str.split(':')
    h = int(parts[0])
    m = int(parts[1]) if len(parts) > 1 else 0
    adjusted_h = h + 24 if h < 5 else h
    return adjusted_h * 60 + m

def process_standard_sheet(ws, day_type: str, direction: str, database: dict, all_stations_set: set):
    """
    일반 본선 시트를 읽어 database[stn][key] 리스트에 추가합니다.
    """
    group_key = f"경춘선_{day_type}_{direction}"

    # 1. Col 1에서 역명 행 매핑 (Row 4부터)
    station_rows = []
    for r in range(4, ws.max_row + 1):
        val = ws.cell(r, 1).value
        if val and str(val).strip():
            clean_str = str(val).strip()
            if clean_str in ('연계열번', '비고', '열차번호', '시발역', '종착역'):
                continue
            stn = normalize_station_name(clean_str)
            station_rows.append((r, stn))
            all_stations_set.add(stn)
            if stn not in database:
                database[stn] = {}
            if group_key not in database[stn]:
                database[stn][group_key] = []

    # 2. 열차별(Col 2부터) 순회
    train_count = 0
    for c in range(2, ws.max_column + 1):
        dest_stn_raw = ws.cell(2, c).value
        train_no_raw = ws.cell(3, c).value

        if not train_no_raw:
            continue

        train_no = str(train_no_raw).strip()
        dest_stn = normalize_station_name(str(dest_stn_raw)) if dest_stn_raw else ''
        is_express = is_express_train(train_no)
        train_count += 1

        for idx, (r, stn) in enumerate(station_rows):
            next_r = station_rows[idx + 1][0] if idx + 1 < len(station_rows) else ws.max_row + 1

            times = []
            for sub_r in range(r, next_r):
                cell_val = ws.cell(sub_r, c).value
                t_str = time_to_hhmm(cell_val)
                if t_str:
                    times.append(t_str)

            if not times:
                continue

            effective_time = times[-1]

            database[stn][group_key].append({
                'no': train_no,
                't': effective_time,
                'd': dest_stn,
                'e': is_express,
            })

    print(f"[{ws.title}] {train_count} trains processed for key: {group_key}")

def process_gwangwoon_sheet(ws, database: dict, all_stations_set: set):
    """
    '광운대시종착(평일)' 시트 처리 (좌측: 상행, 우측: 하행)
    """
    # 1. 좌측 (상행: 춘천 -> 광운대) - Col 1(역명), Col 2~3(열차)
    up_station_rows = []
    for r in range(4, ws.max_row + 1):
        val = ws.cell(r, 1).value
        if val and str(val).strip():
            clean_str = str(val).strip()
            if clean_str in ('연계열번', '비고', '열차번호', '시발역', '종착역'):
                continue
            stn = normalize_station_name(clean_str)
            up_station_rows.append((r, stn))
            all_stations_set.add(stn)
            if stn not in database:
                database[stn] = {}
            if '경춘선_DAY_UP' not in database[stn]:
                database[stn]['경춘선_DAY_UP'] = []

    for c in [2, 3]:
        dest_stn_raw = ws.cell(2, c).value
        train_no_raw = ws.cell(3, c).value
        if not train_no_raw:
            continue
        train_no = str(train_no_raw).strip()
        dest_stn = normalize_station_name(str(dest_stn_raw)) if dest_stn_raw else '광운대'

        for idx, (r, stn) in enumerate(up_station_rows):
            next_r = up_station_rows[idx + 1][0] if idx + 1 < len(up_station_rows) else ws.max_row + 1
            times = []
            for sub_r in range(r, next_r):
                t_str = time_to_hhmm(ws.cell(sub_r, c).value)
                if t_str:
                    times.append(t_str)
            if not times:
                continue
            database[stn]['경춘선_DAY_UP'].append({
                'no': train_no,
                't': times[-1],
                'd': dest_stn,
                'e': 0,
            })

    # 2. 우측 (하행: 광운대 -> 춘천) - Col 5(역명), Col 6~7(열차)
    down_station_rows = []
    for r in range(4, ws.max_row + 1):
        val = ws.cell(r, 5).value
        if val and str(val).strip():
            clean_str = str(val).strip()
            if clean_str in ('연계열번', '비고', '열차번호', '시발역', '종착역'):
                continue
            stn = normalize_station_name(clean_str)
            down_station_rows.append((r, stn))
            all_stations_set.add(stn)
            if stn not in database:
                database[stn] = {}
            if '경춘선_DAY_DOWN' not in database[stn]:
                database[stn]['경춘선_DAY_DOWN'] = []

    for c in [6, 7]:
        dest_stn_raw = ws.cell(2, c).value
        train_no_raw = ws.cell(3, c).value
        if not train_no_raw:
            continue
        train_no = str(train_no_raw).strip()
        dest_stn = normalize_station_name(str(dest_stn_raw)) if dest_stn_raw else '춘천'

        for idx, (r, stn) in enumerate(down_station_rows):
            next_r = down_station_rows[idx + 1][0] if idx + 1 < len(down_station_rows) else ws.max_row + 1
            times = []
            for sub_r in range(r, next_r):
                t_str = time_to_hhmm(ws.cell(sub_r, c).value)
                if t_str:
                    times.append(t_str)
            if not times:
                continue
            database[stn]['경춘선_DAY_DOWN'].append({
                'no': train_no,
                't': times[-1],
                'd': dest_stn,
                'e': 0,
            })

    print(f"[{ws.title}] 4 shuttle trains processed (2 UP, 2 DOWN)")

def main():
    root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    excel_path = os.path.join(root_dir, 'data', '경춘선_열차시간표.xlsx')
    output_dir = os.path.join(root_dir, 'src', 'data', 'subway', 'timetables')
    gz_output_path = os.path.join(output_dir, 'gyeongchun_subway_timetables.json.gz')
    json_output_path = os.path.join(output_dir, 'gyeongchun_subway_timetables.json')

    print(f"Reading Excel: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    database = {}
    all_stations_set = set()

    # 1. 본선 4개 시트 처리
    sheet_configs = [
        ('경춘선 상행(평일)', 'DAY', 'UP'),
        ('경춘선 하행(평일)', 'DAY', 'DOWN'),
        ('경춘선 상행(휴일)', 'END', 'UP'),
        ('경춘선 하행(휴일)', 'END', 'DOWN'),
    ]

    for sheet_name, day_type, direction in sheet_configs:
        if sheet_name in wb.sheetnames:
            process_standard_sheet(wb[sheet_name], day_type, direction, database, all_stations_set)
        else:
            print(f"WARNING: Sheet '{sheet_name}' not found!")

    # 2. 광운대 셔틀 시트 처리
    if '광운대시종착(평일)' in wb.sheetnames:
        process_gwangwoon_sheet(wb['광운대시종착(평일)'], database, all_stations_set)

    # 3. 토요일(SAT) 키를 휴일(END)과 동일하게 복제
    for stn, group in database.items():
        if '경춘선_END_UP' in group and '경춘선_SAT_UP' not in group:
            group['경춘선_SAT_UP'] = list(group['경춘선_END_UP'])
        if '경춘선_END_DOWN' in group and '경춘선_SAT_DOWN' not in group:
            group['경춘선_SAT_DOWN'] = list(group['경춘선_END_DOWN'])

    # 4. 운행 시각 기준 정렬 및 중복 제거
    total_entries = 0
    for stn, group in database.items():
        for gkey, items in group.items():
            unique_dict = {}
            for item in items:
                k = f"{item['no']}_{item['t']}"
                if k not in unique_dict:
                    unique_dict[k] = item
            sorted_items = sorted(
                unique_dict.values(),
                key=lambda x: to_operational_minutes(x['t'])
            )
            group[gkey] = sorted_items
            total_entries += len(sorted_items)

    final_db = {
        'meta': {
            'updatedAt': datetime.datetime.now().isoformat(),
            'source': '한국철도공사 경춘선 열차시간표',
            'line': '경춘선',
            'subwayId': '1067',
            'totalStations': len(all_stations_set),
            'totalTimetableEntries': total_entries,
            'stations': sorted(list(all_stations_set)),
        },
        'data': database,
    }

    print(f"\nTotal Stations: {len(all_stations_set)}")
    print(f"Total Timetable Entries: {total_entries}")

    # JSON 저장
    os.makedirs(output_dir, exist_ok=True)
    json_bytes = json.dumps(final_db, ensure_ascii=False, indent=2).encode('utf-8')
    with open(json_output_path, 'wb') as f:
        f.write(json_bytes)
    print(f"JSON saved: {json_output_path} ({len(json_bytes) / 1024:.1f} KB)")

    # GZIP 압축 저장
    gz_bytes = gzip.compress(json_bytes, compresslevel=9)
    with open(gz_output_path, 'wb') as f:
        f.write(gz_bytes)
    print(f"GZ saved: {gz_output_path} ({len(gz_bytes) / 1024:.1f} KB)")

if __name__ == '__main__':
    main()
